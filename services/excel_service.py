import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

def generate_records_excel(records):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Blood Bank Records"

    # Header styling
    header_fill = PatternFill(start_color="1A365D", end_color="1A365D", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)
    thin_border = Border(
        left=Side(style='thin', color='CBD5E0'),
        right=Side(style='thin', color='CBD5E0'),
        top=Side(style='thin', color='CBD5E0'),
        bottom=Side(style='thin', color='CBD5E0')
    )

    headers = [
        "Record Number", "Date of Donation", "Donor Name", "Gender", "Age",
        "Mobile Number", "City / District", "Blood Group", "Donation Type",
        "Weight (kg)", "BP (Systolic/Diastolic)", "Hemoglobin (g/dL)",
        "Screening Result", "Blood Bag Number", "Bag Type", "Volume (mL)",
        "Status", "Created By", "Signed By (MO)"
    ]

    ws.append(headers)
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_idx, r in enumerate(records, start=2):
        bp_val = f"{r.get('bp_systolic', '')}/{r.get('bp_diastolic', '')}" if r.get('bp_systolic') else ""
        row_data = [
            r.get('record_number', ''),
            r.get('donation_date', ''),
            r.get('full_name', ''),
            r.get('gender', ''),
            r.get('age', ''),
            r.get('mobile_number', ''),
            r.get('city_district', ''),
            r.get('blood_group_known', ''),
            r.get('donation_type', ''),
            r.get('weight_kg', ''),
            bp_val,
            r.get('hemoglobin_g_dl', ''),
            r.get('screening_outcome', ''),
            r.get('blood_bag_number', ''),
            r.get('bag_type', ''),
            r.get('volume_ml', ''),
            r.get('status', ''),
            r.get('created_by_name', ''),
            r.get('signed_by_name', '')
        ]
        ws.append(row_data)

        for col_idx in range(1, len(headers) + 1):
            c = ws.cell(row=row_idx, column=col_idx)
            c.font = data_font
            c.border = thin_border
            if col_idx in [1, 2, 4, 5, 6, 8, 9, 10, 11, 12, 13, 14, 15, 16, 17]:
                c.alignment = Alignment(horizontal="center", vertical="center")
            else:
                c.alignment = Alignment(horizontal="left", vertical="center")

    # Set column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    ws.row_dimensions[1].height = 28

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
